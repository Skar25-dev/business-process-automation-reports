from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import settings
from app.services.report_service import ReportService, get_columns_for_model


def _rows(path):
    """Devuelve (cabeceras, filas) de la primera hoja de un .xlsx."""
    sheet = load_workbook(path).active
    values = list(sheet.iter_rows(values_only=True))
    return list(values[0]), [list(r) for r in values[1:]]


def test_get_columns_for_model_excludes_id():
    columns = get_columns_for_model("ventas")
    assert "amount" in columns
    assert "id" not in columns


def test_get_columns_for_model_unknown_type_is_empty():
    assert get_columns_for_model("no_existe") == []


def test_unknown_report_type_raises(db_session):
    with pytest.raises(ValueError):
        ReportService(db_session).get_report(report_type="no_existe")


def test_excel_report_has_expected_headers_and_default_30_day_window(db_session):
    path = ReportService(db_session).get_report(report_type="ventas")

    assert Path(path).exists() and str(path).endswith(".xlsx")
    headers, rows = _rows(path)
    assert headers == ["PRODUCT NAME", "CATEGORY", "AMOUNT", "DATE"]
    # 3 ventas recientes; la de hace 40 días queda fuera
    assert len(rows) == 3
    assert "Televisor" not in [r[0] for r in rows]


def test_days_parameter_widens_the_window(db_session):
    path = ReportService(db_session).get_report(report_type="ventas", days=60)
    _, rows = _rows(path)
    assert len(rows) == 4


def test_category_filter(db_session):
    path = ReportService(db_session).get_report(report_type="ventas", category="Hogar")
    _, rows = _rows(path)
    assert len(rows) == 2
    assert {r[1] for r in rows} == {"Hogar"}


def test_category_todas_does_not_filter(db_session):
    path = ReportService(db_session).get_report(report_type="ventas", category="Todas")
    _, rows = _rows(path)
    assert len(rows) == 3


def test_no_data_returns_none(db_session):
    assert ReportService(db_session).get_report(report_type="ventas", category="Inexistente") is None


def test_inventory_report_uses_restock_date(db_session, monkeypatch):
    monkeypatch.setattr(
        settings,
        "report_settings",
        {
            "default_days": 30,
            "included_columns": ["product_name", "stock_quantity", "unit_cost"],
            "sheet_name": "Inventario",
        },
    )
    path = ReportService(db_session).get_report(report_type="inventario")
    headers, rows = _rows(path)
    assert headers == ["PRODUCT NAME", "STOCK QUANTITY", "UNIT COST"]
    assert len(rows) == 2


def test_pdf_report_is_generated(db_session, monkeypatch):
    monkeypatch.setattr(settings, "report_format", "pdf")
    path = ReportService(db_session).get_report(report_type="ventas")

    assert str(path).endswith(".pdf")
    assert Path(path).read_bytes().startswith(b"%PDF")
